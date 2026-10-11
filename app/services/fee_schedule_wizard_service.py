"""New Fee Schedule Wizard support (FSW-1/2/3/8/9).

Setup -> Fee Schedule Setup -> **Add New Schedule** runs the legacy four-step
wizard: schedule type -> offices -> providers / specialties -> fees. Before this
module the frontend applied it as ``POST /fee-schedules`` + ``PUT .../entries/bulk``
+ one ``POST /fee-schedule-assignments`` per provider x office, so a failure on
call 37 of 250 left a schedule half-linked with nothing to roll it back.

* :func:`bulk_assign` (FSW-2) — many assignment rows in one transaction with an
  ``on_conflict`` policy (``error`` | ``rebind`` | ``skip``) and a per-row outcome.
  ``dry_run`` is the conflict probe the wizard's "Currently Bound To" column needs
  (FSW-7) without paging the whole table.
* :func:`create_wizard` (FSW-1) — schedule + entries (+ Copy From Existing) +
  assignments + office default pointers, **all or nothing**.
* :func:`scope` (FSW-9) — the inverse: what a schedule is bound to, shaped so
  "Edit Schedule" can re-enter the wizard.
* :func:`import_entries` (FSW-8) — the "Import From Excel File" template as a
  server import (``.csv`` or native ``.xlsx``) with a row-level report.

**FSW-3 — Office + All Providers is the office pointer, by design.** Denticon
leaves ``FeeScheH.OID`` empty on all 36 exported headers and its ``FeeScheA``
holds no office-only row: an "Office" schedule applied to specific offices for
all providers *is* ``Office.PATIENTFEEID`` — ``offices.default_fee_schedule_id``,
tier 5 of the precedence card. An office-scoped assignment row would be a second,
competing way to say the same thing (the eight all-NULL migrated rows are what
that looks like), so scope-only rows stay refused. Because the pointer is single
valued, the wizard response reports the pointer each office held before
(``previous_default_fee_schedule_id``) so replacing one is never silent.

Every writer here is **non-committing until the end** and reuses the guards the
single-row CRUD paths use (``fee_schedule_service``), so the bulk paths cannot
accept a row the per-row path would refuse.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.exceptions import AppError, ConflictError, ValidationError
from app.db.models import (
    FeeSchedule,
    FeeScheduleAssignment,
    FeeScheduleEntry,
    InsuranceCarrier,
    InsurancePlan,
    Office,
    OfficeGroup,
    Patient,
    ProcedureCode,
    Provider,
)
from app.services import fee_vocab, provider_taxonomy_service
from app.services import fee_schedule_service as fs_svc

ON_CONFLICT_POLICIES: tuple[str, ...] = ("error", "rebind", "skip")
#: Per-call cap. 50 providers x 5 offices is the report's worked example (250).
MAX_ASSIGNMENT_ROWS = 2000
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 20000

_ASSIGNMENT_FIELDS = fee_vocab.ASSIGNMENT_KEYS + ("fee_schedule_id",)


def _err(code: str) -> str:
    return fee_vocab.ERROR_CODES.get(code, code)


def _clean(value):  # noqa: ANN001, ANN202
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


# ── FSW-2: bulk assignments ──────────────────────────────────────────────────


def _in_tenant(db: Session, model, obj_id, tenant_id: int) -> object | None:  # noqa: ANN001
    obj = db.get(model, obj_id)
    if obj is None or getattr(obj, "tenant_id", tenant_id) != tenant_id:
        return None
    return obj


def _check_refs(db: Session, row: dict, tenant_id: int) -> tuple[str, str] | None:
    """The first foreign key on ``row`` that does not resolve inside the tenant,
    as ``(error_code, field)``."""
    checks = (
        ("provider_id", Provider, "assignment_provider_invalid"),
        ("office_id", Office, "assignment_office_invalid"),
        ("office_group_id", OfficeGroup, "assignment_office_invalid"),
        ("ins_plan_id", InsurancePlan, "assignment_target_invalid"),
        ("carrier_id", InsuranceCarrier, "assignment_target_invalid"),
    )
    for field, model, code in checks:
        value = row.get(field)
        if value is not None and _in_tenant(db, model, value, tenant_id) is None:
            return code, field
    return None


def _schedule_names(db: Session, ids: set[int]) -> dict[int, str]:
    ids = {i for i in ids if i is not None}
    if not ids:
        return {}
    return dict(db.execute(select(FeeSchedule.id, FeeSchedule.name).where(FeeSchedule.id.in_(ids))).all())


def apply_assignments(
    db: Session,
    tenant_id: int,
    rows: list[dict],
    *,
    on_conflict: str = "error",
    user_id: int | None = None,
    fee_schedule_id: int | None = None,
    write: bool = True,
) -> tuple[list[dict], bool]:
    """Validate and (when ``write``) stage every row. Never commits.

    Returns ``(outcomes, ok)``. Each outcome is ``{index, status, ...}`` with
    status ``created`` | ``rebound`` | ``unchanged`` | ``skipped`` | ``conflict`` |
    ``error``. ``ok`` is false when any row is ``conflict`` or ``error`` — the
    caller then rolls the whole batch back. ``fee_schedule_id`` (the wizard's new
    schedule) overrides every row's own.

    A row already bound to the *same* schedule is ``unchanged`` under every
    policy, so re-applying a wizard is idempotent rather than a wall of 409s.
    """
    if on_conflict not in ON_CONFLICT_POLICIES:
        raise ValidationError("on_conflict must be one of error, rebind, skip",
                              details={"code": "invalid_on_conflict", "field": "on_conflict",
                                       "allowed": list(ON_CONFLICT_POLICIES)})
    if len(rows) > MAX_ASSIGNMENT_ROWS:
        raise ValidationError(f"At most {MAX_ASSIGNMENT_ROWS} assignment rows per request",
                              details={"code": "too_many_rows", "max": MAX_ASSIGNMENT_ROWS})

    outcomes: list[dict] = []
    seen: dict[tuple, int] = {}
    valid_schedules: dict[int, bool] = {}
    ok = True
    for index, raw in enumerate(rows):
        row = {k: _clean(raw.get(k)) for k in _ASSIGNMENT_FIELDS}
        if fee_schedule_id is not None:
            row["fee_schedule_id"] = fee_schedule_id
        outcome: dict = {"index": index, "keys": {k: row[k] for k in fee_vocab.ASSIGNMENT_KEYS},
                         "fee_schedule_id": row["fee_schedule_id"]}
        outcomes.append(outcome)

        def fail(code: str, field: str | None = None, *, _o=outcome) -> None:
            _o.update(status="error", error_code=code, message=_err(code))
            if field:
                _o["field"] = field

        try:
            fs_svc.normalise_assignment_specialty(row)
        except ValidationError:
            fail("assignment_specialty_unknown", "specialty_id")
            ok = False
            continue
        outcome["keys"]["specialty_id"] = row["specialty_id"]

        if not fee_vocab.has_assignment_target(row):
            fail("assignment_needs_target")
            ok = False
            continue
        sched_id = row["fee_schedule_id"]
        if sched_id not in valid_schedules:
            valid_schedules[sched_id] = (
                fee_schedule_id is not None  # the wizard's own, freshly created schedule
                or fs_svc._schedule_valid(db, sched_id, tenant_id)
            )
        if not valid_schedules[sched_id]:
            fail("assignment_schedule_invalid", "fee_schedule_id")
            ok = False
            continue
        bad_ref = _check_refs(db, row, tenant_id)
        if bad_ref is not None:
            fail(*bad_ref)
            ok = False
            continue
        key = tuple(row[k] for k in fee_vocab.ASSIGNMENT_KEYS)
        if key in seen:
            fail("assignment_duplicate_in_request")
            outcome["duplicate_of_index"] = seen[key]
            ok = False
            continue
        seen[key] = index

        existing = fs_svc.find_assignment(db, row, tenant_id)
        if existing is None:
            outcome["status"] = "created"
            if write:
                created = FeeScheduleAssignment(
                    tenant_id=tenant_id, fee_schedule_id=sched_id,
                    **{k: row[k] for k in fee_vocab.ASSIGNMENT_KEYS},
                    created_by=(str(user_id) if user_id is not None else None),
                )
                db.add(created)
                db.flush()
                outcome["assignment_id"] = created.id
            continue

        outcome["assignment_id"] = existing.id
        outcome["previous_fee_schedule_id"] = existing.fee_schedule_id
        if existing.fee_schedule_id == sched_id:
            outcome["status"] = "unchanged"
        elif on_conflict == "rebind":
            outcome["status"] = "rebound"
            if write:
                existing.fee_schedule_id = sched_id
                existing.updated_by = user_id
        elif on_conflict == "skip":
            outcome["status"] = "skipped"
        else:
            outcome.update(status="conflict", error_code="assignment_duplicate_target",
                           message=_err("assignment_duplicate_target"))
            ok = False

    names = _schedule_names(db, {o.get("previous_fee_schedule_id") for o in outcomes})
    for o in outcomes:
        if o.get("previous_fee_schedule_id") is not None:
            o["previous_fee_schedule_name"] = names.get(o["previous_fee_schedule_id"])
    if write:
        db.flush()
    return outcomes, ok


def _summary(outcomes: list[dict]) -> dict:
    counts = {s: 0 for s in ("created", "rebound", "unchanged", "skipped", "conflict", "error")}
    for o in outcomes:
        counts[o["status"]] = counts.get(o["status"], 0) + 1
    return counts


def _failure(code: str, message: str, report: dict, outcomes: list[dict]) -> AppError:
    """422 when any row is invalid, else 409 (only conflicts). The full report rides
    in ``details`` so the client can render every row's outcome."""
    has_error = any(o["status"] == "error" for o in outcomes)
    cls = ValidationError if has_error else ConflictError
    return cls(message, details={"code": code, **report})


def bulk_assign(
    db: Session,
    tenant_id: int,
    rows: list[dict],
    *,
    on_conflict: str = "error",
    dry_run: bool = False,
    user_id: int | None = None,
) -> dict:
    """``POST /fee-schedule-assignments/bulk`` (FSW-2). All rows or none: a single
    invalid row (or, under ``on_conflict="error"``, a single conflict) rolls the
    batch back and answers 422/409 with every row's outcome. ``dry_run`` answers
    200 with the same report and writes nothing."""
    outcomes, ok = apply_assignments(db, tenant_id, rows, on_conflict=on_conflict,
                                     user_id=user_id, write=not dry_run)
    report = {"ok": ok, "dry_run": dry_run, "on_conflict": on_conflict,
              "summary": _summary(outcomes), "outcomes": outcomes}
    if dry_run:
        db.rollback()
        return report
    if not ok:
        db.rollback()
        raise _failure("bulk_assignment_failed", "No assignment was written; see details.outcomes",
                       report, outcomes)
    db.commit()
    return report


# ── FSW-1 (+FSW-3): the atomic wizard ────────────────────────────────────────


def _known_codes(db: Session, codes: set[str]) -> set[str]:
    if not codes:
        return set()
    return set(db.execute(select(ProcedureCode.code).where(ProcedureCode.code.in_(codes))).scalars().all())


def _latest_entries(db: Session, schedule_id: int, on: date | None) -> dict[str, FeeScheduleEntry]:
    """Each code's entry in force on ``on`` (latest ``effective_date`` <= ``on``),
    else its earliest future one — Copy From Existing copies today's prices."""
    rows = db.execute(
        select(FeeScheduleEntry).where(FeeScheduleEntry.fee_schedule_id == schedule_id)
    ).scalars().all()
    best: dict[str, FeeScheduleEntry] = {}
    for row in rows:
        eff = row.effective_date or date.min
        cur = best.get(row.procedure_code)
        if cur is None:
            best[row.procedure_code] = row
            continue
        cur_eff = cur.effective_date or date.min
        if on is None:
            if eff > cur_eff:
                best[row.procedure_code] = row
        elif eff <= on and (cur_eff > on or eff > cur_eff):
            best[row.procedure_code] = row
        elif eff > on and cur_eff > on and eff < cur_eff:
            best[row.procedure_code] = row
    return best


def _schedule_payload(schedule: dict) -> dict:
    allowed = {"name", "fee_type", "pricing_model", "effective_date", "legacy_id"}
    data = {k: v for k, v in schedule.items() if k in allowed}
    if not _clean(data.get("name")):
        raise ValidationError("A fee schedule needs a name",
                              details={"code": "fee_schedule_name_required", "field": "name",
                                       "step": "schedule"})
    data["name"] = data["name"].strip()
    fs_svc.FeeScheduleCRUD._normalise(data)
    fs_svc.FeeScheduleCRUD._guard_pricing_model(data, current=None)
    data.setdefault("fee_type", fee_vocab.DEFAULT_FEE_TYPE)
    data["fee_type"] = data["fee_type"] or fee_vocab.DEFAULT_FEE_TYPE
    data["pricing_model"] = data.get("pricing_model") or fee_vocab.DEFAULT_PRICING_MODEL
    return data


def create_wizard(db: Session, tenant_id: int, payload: dict, *, user_id: int | None = None) -> dict:
    """``POST /fee-schedules/wizard`` (FSW-1) — one transaction for the whole
    wizard. Any invalid entry / assignment / office, or an assignment conflict
    under ``on_conflict="error"``, rolls everything back (422 / 409) and returns
    the full per-row report in ``details``. ``dry_run`` validates and reports
    without writing (``fee_schedule.id`` is then ``null``)."""
    dry_run = bool(payload.get("dry_run"))
    on_conflict = payload.get("on_conflict") or "error"
    try:
        sched_data = _schedule_payload(dict(payload.get("schedule") or {}))
    except ValidationError as exc:
        exc.details = {**(exc.details or {}), "step": "schedule"}
        raise

    sched = FeeSchedule(tenant_id=tenant_id, is_active=True, version=1, **sched_data)
    db.add(sched)
    db.flush()
    # Step 1's effective date is the date the wizard's prices start.
    eff = payload.get("entries_effective_date") or sched.effective_date or date.today()

    # ── step 4: entries (copy-from first, explicit rows override) ────────────
    entry_errors: list[dict] = []
    staged: dict[str, dict] = {}
    copied = dropped_plan_pays = 0
    source_id = payload.get("copy_from_schedule_id")
    if source_id is not None:
        # A retired list is still a fine thing to copy prices from.
        if _in_tenant(db, FeeSchedule, source_id, tenant_id) is None:
            db.rollback()
            raise ValidationError("The schedule to copy from was not found in this tenant",
                                  details={"code": "copy_source_invalid", "field": "copy_from_schedule_id",
                                           "step": "entries"})
        allows = fee_vocab.allows_plan_pays(sched.fee_type, sched.pricing_model)
        for code, src in _latest_entries(db, source_id, eff).items():
            insurance_fee = src.insurance_fee
            if insurance_fee is not None and insurance_fee > 0 and not allows:
                insurance_fee = None
                dropped_plan_pays += 1
            staged[code] = {"patient_fee": src.patient_fee, "insurance_fee": insurance_fee,
                            "amb_code": src.amb_code, "is_no_charge": src.is_no_charge}
            copied += 1

    entries = payload.get("entries") or []
    known = _known_codes(db, {(e.get("procedure_code") or "").strip().upper() for e in entries} - {""})
    for index, item in enumerate(entries):
        code = (item.get("procedure_code") or "").strip().upper()
        if not code:
            entry_errors.append({"index": index, "error_code": "procedure_code_required",
                                 "message": "procedure_code is required"})
            continue
        if code not in known:
            entry_errors.append({"index": index, "procedure_code": code,
                                 "error_code": "unknown_procedure_code",
                                 "message": _err("unknown_procedure_code")})
            continue
        try:
            fs_svc._validate_entry_amounts(sched, item.get("patient_fee"), item.get("insurance_fee"))
        except ValidationError as exc:
            entry_errors.append({"index": index, "procedure_code": code,
                                 "error_code": exc.details.get("code"), "field": exc.details.get("field"),
                                 "message": exc.message})
            continue
        fields = {k: item[k] for k in ("patient_fee", "insurance_fee", "amb_code", "is_no_charge")
                  if k in item and item[k] is not None}
        staged.setdefault(code, {}).update(fields)
        if item.get("effective_date"):
            staged[code]["_eff"] = item["effective_date"]

    created = 0
    for code, fields in staged.items():
        if fs_svc._upsert_entry(db, sched.id, tenant_id, code, fields.pop("_eff", None) or eff, fields):
            created += 1
    db.flush()

    # ── step 3: assignments ──────────────────────────────────────────────────
    outcomes, assignments_ok = apply_assignments(
        db, tenant_id, payload.get("assignments") or [], on_conflict=on_conflict,
        user_id=user_id, fee_schedule_id=sched.id,
    )

    # ── FSW-3: office default pointers ───────────────────────────────────────
    office_results: list[dict] = []
    office_errors: list[dict] = []
    defaults = payload.get("office_defaults") or {}
    set_default = bool(defaults.get("set_default_fee_schedule"))
    set_ucr = bool(defaults.get("set_ucr_fee_schedule"))
    for office_id in defaults.get("office_ids") or []:
        office = _in_tenant(db, Office, office_id, tenant_id)
        if office is None:
            office_errors.append({"office_id": office_id, "error_code": "assignment_office_invalid",
                                  "message": _err("assignment_office_invalid")})
            continue
        result = {"office_id": office.id, "office_name": office.name}
        if set_default:
            result["previous_default_fee_schedule_id"] = office.default_fee_schedule_id
            result["replaced_default"] = office.default_fee_schedule_id not in (None, sched.id)
            office.default_fee_schedule_id = sched.id
            result["default_fee_schedule_id"] = sched.id
        if set_ucr:
            result["previous_default_ucr_fee_schedule_id"] = office.default_ucr_fee_schedule_id
            result["replaced_ucr"] = office.default_ucr_fee_schedule_id not in (None, sched.id)
            office.default_ucr_fee_schedule_id = sched.id
            result["default_ucr_fee_schedule_id"] = sched.id
        office_results.append(result)
    prev_names = _schedule_names(db, {r.get("previous_default_fee_schedule_id") for r in office_results}
                                 | {r.get("previous_default_ucr_fee_schedule_id") for r in office_results})
    for r in office_results:
        for k in ("previous_default_fee_schedule_id", "previous_default_ucr_fee_schedule_id"):
            if r.get(k) is not None:
                r[k.replace("_id", "_name")] = prev_names.get(r[k])

    ok = assignments_ok and not entry_errors and not office_errors
    report = {
        "ok": ok,
        "dry_run": dry_run,
        "fee_schedule": {
            "id": None if dry_run else sched.id, "name": sched.name, "fee_type": sched.fee_type,
            "pricing_model": sched.pricing_model,
            "effective_date": sched.effective_date.isoformat() if sched.effective_date else None,
            "is_active": True,
        },
        "entries": {"staged": len(staged), "created": created, "copied": copied,
                    "dropped_plan_pays": dropped_plan_pays,
                    "effective_date": eff.isoformat(),
                    "errors": entry_errors},
        "assignments": {"on_conflict": on_conflict, "summary": _summary(outcomes), "outcomes": outcomes},
        "office_defaults": {"results": office_results, "errors": office_errors},
    }
    if dry_run:
        db.rollback()
        return report
    if not ok:
        db.rollback()
        has_error = bool(entry_errors or office_errors) or any(o["status"] == "error" for o in outcomes)
        cls = ValidationError if has_error else ConflictError
        raise cls("The fee schedule was not created; see details", details={"code": "fee_schedule_wizard_failed",
                                                                              **report})
    db.commit()
    return report


# ── FSW-9: the scope read ────────────────────────────────────────────────────


def scope(db: Session, schedule_id: int, tenant_id: int) -> dict:
    """``GET /fee-schedules/{id}/scope`` — every place this schedule is bound:
    assignment rows (denormalised), the offices whose default / UCR pointer names
    it, the patients on it, and a ``wizard`` block in the wizard's own vocabulary
    (offices all|specific, providers all|providers|specialties) so "Edit Schedule"
    can re-enter the wizard."""
    sched = fs_svc._get_in_tenant(db, schedule_id, tenant_id)
    rows = db.execute(
        select(FeeScheduleAssignment).where(
            FeeScheduleAssignment.tenant_id == tenant_id,
            FeeScheduleAssignment.fee_schedule_id == sched.id,
        ).order_by(FeeScheduleAssignment.id)
    ).scalars().all()

    def _names(model, ids, attr="name") -> dict:  # noqa: ANN001
        ids = {i for i in ids if i is not None}
        if not ids:
            return {}
        return dict(db.execute(select(model.id, getattr(model, attr)).where(model.id.in_(ids))).all())

    offices = _names(Office, {r.office_id for r in rows})
    providers = _names(Provider, {r.provider_id for r in rows})
    carriers = _names(InsuranceCarrier, {r.carrier_id for r in rows})
    plans = _names(InsurancePlan, {r.ins_plan_id for r in rows}, "group_number")
    groups = _names(OfficeGroup, {r.office_group_id for r in rows})

    assignments = []
    for r in rows:
        spec = provider_taxonomy_service.canonical_specialty(r.specialty_id)
        assignments.append({
            "id": r.id, "ins_plan_id": r.ins_plan_id, "ins_plan_group_number": plans.get(r.ins_plan_id),
            "carrier_id": r.carrier_id, "carrier_name": carriers.get(r.carrier_id),
            "provider_id": r.provider_id, "provider_name": providers.get(r.provider_id),
            "specialty_id": r.specialty_id, "specialty_label": provider_taxonomy_service.label_for(spec),
            "office_id": r.office_id, "office_name": offices.get(r.office_id),
            "office_group_id": r.office_group_id, "office_group_name": groups.get(r.office_group_id),
            "fee_source": fee_vocab.fee_source_for_assignment(r),
        })

    pointer_rows = db.execute(
        select(Office.id, Office.name, Office.default_fee_schedule_id, Office.default_ucr_fee_schedule_id)
        .where(Office.tenant_id == tenant_id,
               (Office.default_fee_schedule_id == sched.id) | (Office.default_ucr_fee_schedule_id == sched.id))
        .order_by(Office.id)
    ).all()
    office_pointers = [
        {"office_id": oid, "office_name": name,
         "is_default_fee_schedule": dflt == sched.id, "is_ucr_fee_schedule": ucr == sched.id}
        for oid, name, dflt, ucr in pointer_rows
    ]
    patients = int(db.execute(select(func.count()).select_from(Patient).where(
        Patient.tenant_id == tenant_id, Patient.fee_schedule_id == sched.id)).scalar() or 0)

    # Provider/specialty rows only: payer rows are a different wizard path, and a
    # legacy all-NULL row (no target at all) binds nobody, so it must not read
    # as "All Offices".
    person_rows = [r for r in rows
                   if fee_vocab.has_assignment_target(r) and not (r.ins_plan_id or r.carrier_id)]
    provider_ids = sorted({r.provider_id for r in person_rows if r.provider_id})
    specialty_ids = sorted({r.specialty_id for r in person_rows if r.specialty_id})
    office_ids = sorted({r.office_id for r in person_rows if r.office_id is not None}
                        | {p["office_id"] for p in office_pointers if p["is_default_fee_schedule"]})
    all_offices = bool(person_rows) and any(r.office_id is None and r.office_group_id is None
                                            for r in person_rows)
    wizard = {
        "offices": {"mode": "all" if (all_offices or not office_ids) else "specific",
                    "office_ids": office_ids},
        "providers": {
            "mode": "providers" if provider_ids else ("specialties" if specialty_ids else "all"),
            "provider_ids": provider_ids, "specialty_ids": specialty_ids,
        },
        "is_office_default": any(p["is_default_fee_schedule"] for p in office_pointers),
    }
    return {
        "fee_schedule_id": sched.id, "name": sched.name, "fee_type": sched.fee_type,
        "pricing_model": sched.pricing_model, "is_active": bool(sched.is_active),
        "assignments": assignments,
        "offices": [{"office_id": k, "office_name": v} for k, v in sorted(offices.items())],
        "providers": [{"provider_id": k, "provider_name": v} for k, v in sorted(providers.items())],
        "specialties": [{"specialty_id": s, "label": provider_taxonomy_service.label_for(
            provider_taxonomy_service.canonical_specialty(s))} for s in specialty_ids],
        "carriers": [{"carrier_id": k, "carrier_name": v} for k, v in sorted(carriers.items())],
        "plans": [{"ins_plan_id": k, "group_number": v} for k, v in sorted(plans.items())],
        "office_pointers": office_pointers,
        "patients": patients,
        "wizard": wizard,
    }


# ── FSW-8: server-side import ────────────────────────────────────────────────

#: Normalised header -> field. The first four are the Utilities template's own
#: columns (``procedure_code, description, fee, effective_date``); the rest are
#: the spellings a practice's own spreadsheet tends to use.
_HEADER_SYNONYMS: dict[str, str] = {
    "procedure_code": "procedure_code", "code": "procedure_code", "proc_code": "procedure_code",
    "procedure": "procedure_code", "ada_code": "procedure_code", "cdt_code": "procedure_code",
    "fee": "patient_fee", "patient_fee": "patient_fee", "amount": "patient_fee", "ucr_fee": "patient_fee",
    "office_fee": "patient_fee", "pat_amt": "patient_fee", "patamt": "patient_fee",
    "insurance_fee": "insurance_fee", "plan_pays": "insurance_fee", "ins_fee": "insurance_fee",
    "ins_amt": "insurance_fee", "insamt": "insurance_fee",
    "amb_code": "amb_code", "amb": "amb_code", "ambcode": "amb_code",
    "effective_date": "effective_date", "eff_date": "effective_date", "effective": "effective_date",
    "no_charge": "is_no_charge", "is_no_charge": "is_no_charge",
    "description": "description", "desc": "description",
}
_DATE_FORMATS = ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y")
_TRUE = {"y", "yes", "true", "1", "x"}


def _norm_header(h) -> str:  # noqa: ANN001
    return re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_")


def _read_rows(filename: str, content: bytes) -> tuple[str, list[list]]:
    is_xlsx = (filename or "").lower().endswith((".xlsx", ".xlsm")) or content[:2] == b"PK"
    if is_xlsx:
        try:
            from openpyxl import load_workbook  # noqa: PLC0415 (lazy: only the import path needs it)
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001 — any unreadable workbook is the same 422
            raise ValidationError("The file is not a readable .xlsx workbook",
                                  details={"code": "import_file_unreadable"}) from exc
        sheet = wb.worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        wb.close()
        return "xlsx", rows
    if (filename or "").lower().endswith(".xls"):
        raise ValidationError("Legacy .xls is not supported; save as .xlsx or .csv",
                              details={"code": "import_file_type_unsupported"})
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    return "csv", [row for row in csv.reader(io.StringIO(text))]


def _parse_money(value) -> Decimal | None:  # noqa: ANN001
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    cleaned = str(value).strip().replace("$", "").replace(",", "")
    if cleaned.startswith("(") and cleaned.endswith(")"):
        cleaned = "-" + cleaned[1:-1]
    return Decimal(cleaned)


def _parse_date(value) -> date | None:  # noqa: ANN001
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def import_entries(
    db: Session,
    schedule_id: int,
    tenant_id: int,
    *,
    filename: str,
    content: bytes,
    effective_date: date | None = None,
    dry_run: bool = False,
    all_or_nothing: bool = False,
) -> dict:
    """``POST /fee-schedules/{id}/entries/import`` (FSW-8). Accepts the Utilities
    template (``procedure_code, description, fee, effective_date``) as ``.csv`` or
    native ``.xlsx``, plus optional ``plan_pays`` / ``amb_code`` / ``no_charge``
    columns, matched by header name rather than position.

    Every row is validated the way the entry CRUD validates one (known code, no
    negative fee, Plan Pays only on a copay list) and reported by its 1-based data
    row number. Valid rows are upserted on ``(schedule, code, effective_date)``;
    invalid ones are skipped — unless ``all_or_nothing``, when any error writes
    nothing. ``dry_run`` is the validation report alone."""
    sched = fs_svc._get_in_tenant(db, schedule_id, tenant_id)
    if len(content) > MAX_IMPORT_BYTES:
        raise ValidationError("The file is too large",
                              details={"code": "import_file_too_large", "max_bytes": MAX_IMPORT_BYTES})
    file_type, rows = _read_rows(filename, content)
    rows = [r for r in rows if any(c not in (None, "") and str(c).strip() for c in r)]
    if not rows:
        raise ValidationError("The file is empty", details={"code": "import_file_empty"})

    header = [_norm_header(h) for h in rows[0]]
    columns: dict[str, int] = {}
    for idx, h in enumerate(header):
        field = _HEADER_SYNONYMS.get(h)
        if field and field not in columns:
            columns[field] = idx
    header_ok = "procedure_code" in columns and bool(
        {"patient_fee", "insurance_fee", "is_no_charge"} & columns.keys())
    if not header_ok:
        raise ValidationError(
            "The header row needs a procedure_code column and a fee column",
            details={"code": "import_header_invalid", "header": [str(h) for h in rows[0]],
                     "expected": ["procedure_code", "description", "fee", "effective_date"]})
    data_rows = rows[1:]
    if len(data_rows) > MAX_IMPORT_ROWS:
        raise ValidationError(f"At most {MAX_IMPORT_ROWS} rows per import",
                              details={"code": "too_many_rows", "max": MAX_IMPORT_ROWS})

    def cell(row: list, field: str):  # noqa: ANN202
        idx = columns.get(field)
        return row[idx] if idx is not None and idx < len(row) else None

    default_eff = effective_date or sched.effective_date or date.today()
    codes = {str(cell(r, "procedure_code") or "").strip().upper() for r in data_rows} - {""}
    known = _known_codes(db, codes)
    errors: list[dict] = []
    valid: list[tuple[int, str, date, dict]] = []
    seen: dict[tuple[str, date], int] = {}
    for n, row in enumerate(data_rows, start=1):
        code = str(cell(row, "procedure_code") or "").strip().upper()

        def bad(error_code: str, message: str, *, _n=n, _code=code) -> None:
            errors.append({"row": _n, "procedure_code": _code or None, "code": error_code, "message": message})

        if not code:
            bad("procedure_code_required", "procedure_code is required")
            continue
        if code not in known:
            bad("unknown_procedure_code", f'"{code}" is not a procedure code')
            continue
        try:
            patient_fee = _parse_money(cell(row, "patient_fee"))
            insurance_fee = _parse_money(cell(row, "insurance_fee"))
        except InvalidOperation:
            bad("invalid_amount", "fee is not a valid amount")
            continue
        try:
            eff = _parse_date(cell(row, "effective_date")) or default_eff
        except ValueError:
            bad("invalid_effective_date", "effective_date must be YYYY-MM-DD (or MM/DD/YYYY)")
            continue
        no_charge_raw = cell(row, "is_no_charge")
        is_no_charge = (str(no_charge_raw).strip().lower() in _TRUE) if no_charge_raw not in (None, "") else None
        if patient_fee is None and insurance_fee is None and not is_no_charge:
            bad("fee_entry_needs_amount", _err("fee_entry_needs_amount"))
            continue
        try:
            fs_svc._validate_entry_amounts(sched, patient_fee, insurance_fee)
        except ValidationError as exc:
            bad(exc.details.get("code") or "invalid_amount", exc.message)
            continue
        if (code, eff) in seen:
            bad("duplicate_row", f"{code} from {eff.isoformat()} already appears on row {seen[(code, eff)]}")
            continue
        seen[(code, eff)] = n
        fields: dict = {}
        if patient_fee is not None:
            fields["patient_fee"] = patient_fee
        if insurance_fee is not None:
            fields["insurance_fee"] = insurance_fee
        amb = cell(row, "amb_code")
        if amb not in (None, "") and str(amb).strip():
            fields["amb_code"] = str(amb).strip().upper()
        if is_no_charge is not None:
            fields["is_no_charge"] = is_no_charge
        valid.append((n, code, eff, fields))

    write = not dry_run and not (all_or_nothing and errors)
    created = updated = 0
    if write:
        for _n, code, eff, fields in valid:
            if fs_svc._upsert_entry(db, sched.id, tenant_id, code, eff, fields):
                created += 1
            else:
                updated += 1
        db.commit()
    return {
        "fee_schedule_id": sched.id,
        "file_type": file_type,
        "header_ok": True,
        "columns": {k: str(rows[0][v]) for k, v in columns.items()},
        "total_rows": len(data_rows),
        "valid_rows": len(valid),
        "error_rows": len(errors),
        "created": created,
        "updated": updated,
        "written": write,
        "dry_run": dry_run,
        "all_or_nothing": all_or_nothing,
        "effective_date": default_eff.isoformat(),
        "errors": errors,
    }


# ── FSW-5: the specialty catalog with live counts ────────────────────────────


def specialty_catalog(db: Session, tenant_id: int, *, include_inactive: bool = False) -> list[dict]:
    """The dental taxonomy catalog with the tenant's providers grouped under each
    code (stored ``taxonomy_code``, else derived from ``specialty`` — never the
    General-Dentist default), plus any taxonomy-shaped code a provider carries
    that the catalog does not list. This is what the wizard's Specialties option
    renders instead of a free-text list that is blank on 95 of 97 providers."""
    stmt = select(Provider).where(Provider.tenant_id == tenant_id)
    if not include_inactive:
        stmt = stmt.where(Provider.is_active.is_(True))
    by_code: dict[str, list[str]] = {}
    unclassified = 0
    for p in db.execute(stmt).scalars().all():
        code = provider_taxonomy_service.provider_specialty_code(p.taxonomy_code, p.specialty)
        if code is None:
            unclassified += 1
            continue
        by_code.setdefault(code, []).append(p.id)
    out = []
    for code, label, _ in provider_taxonomy_service.TAXONOMY_CODES:
        ids = sorted(by_code.pop(code, []))
        out.append({"specialty_id": code, "label": label, "in_catalog": True,
                    "provider_count": len(ids), "provider_ids": ids})
    for code, ids in sorted(by_code.items()):
        out.append({"specialty_id": code, "label": provider_taxonomy_service.label_for(code) or code,
                    "in_catalog": False, "provider_count": len(ids), "provider_ids": sorted(ids)})
    out.append({"specialty_id": None, "label": "No specialty", "in_catalog": False,
                "provider_count": unclassified, "provider_ids": []})
    return out
